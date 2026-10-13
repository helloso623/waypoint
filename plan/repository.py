"""Local file repository: ingest, categorize, TF-IDF RAG, rolling master summary."""
import io
import json
import os
import re
import shutil
import subprocess

import proc_util
import sys
import uuid
import zipfile
from datetime import date as date_cls
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

import ai
import budget
import db
import orchestrator

REPO_STORE = Path(__file__).parent / "repo_store"
REPO_STORE.mkdir(exist_ok=True)
VERSIONS_DIR = REPO_STORE / ".versions"
VERSIONS_DIR.mkdir(exist_ok=True)
MAX_VERSION_BACKUPS = 3
RUN_TIMEOUT_SECONDS = 30
VENV_PYTHON = Path(os.environ.get("WAYPOINT_PYTHON") or sys.executable)
ZIP_TEXT_EXTENSIONS = {
    ".tex", ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json",
    ".py", ".js", ".ts", ".ipynb", ".bib", ".sty", ".cls",
}

CATEGORIES = ("Docs", "Code", "Notes", "Data", "Research", "Media", "Google", "Other")

CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".scala", ".sh", ".bash",
    ".sql", ".html", ".css", ".scss", ".vue", ".r", ".m", ".lua",
}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".tex"}
DATA_EXTENSIONS = {".json", ".yaml", ".yml", ".xml", ".toml", ".ini", ".cfg"}
NOTEBOOK_EXTENSIONS = {".ipynb"}
DOC_EXTENSIONS = {".pdf", ".docx", ".doc", ".rtf", ".odt"}
SHEET_EXTENSIONS = {".csv", ".tsv", ".xls", ".xlsx", ".ods", ".numbers"}
SLIDE_EXTENSIONS = {".ppt", ".pptx", ".odp", ".key"}
MEDIA_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".mp3", ".mp4", ".wav"}
WORKSPACE_KINDS = ("docs", "sheets", "slides", "other")

CHUNK_SIZE = 4000
CHUNK_OVERLAP = 400
SNIPPET_RADIUS = 120
AI_MIN_WORDS = 40
AI_PRESSURE_MAX = 0.85
# Full extracted text for content/API consumers (list endpoints still use a short preview).
EXTRACTED_TEXT_RETURN_MAX = 500_000
ASK_DOCS_CONTEXT_CHARS = 48_000
ASK_DOCS_MAX_TOKENS = 1200

_rag_cache = {"file_count": None, "vectorizer": None, "matrix": None, "chunks": []}


def _word_count(text):
    return len(re.findall(r"\S+", text or ""))


def _extract_pdf_text(stream):
    reader = PdfReader(stream)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(stream):
    document = Document(stream)
    return "\n".join(p.text for p in document.paragraphs)


def _extract_ipynb_text(stream):
    raw = stream.read()
    data = json.loads(raw.decode("utf-8", errors="ignore"))
    parts = []
    for cell in data.get("cells", []):
        src = cell.get("source", [])
        if isinstance(src, list):
            src = "".join(src)
        cell_type = cell.get("cell_type", "code")
        if src.strip():
            parts.append(f"[{cell_type}]\n{src}")
    return "\n\n".join(parts).strip(), None


def _extract_from_bytes(filename, data):
    """Return (text, note) from raw bytes (used after content edits and Google import)."""
    suffix = Path(filename).suffix.lower()
    stream = io.BytesIO(data)
    if suffix == ".pdf":
        return _extract_pdf_text(stream).strip(), None
    if suffix in (".docx", ".doc"):
        return _extract_docx_text(stream).strip(), None
    if suffix in NOTEBOOK_EXTENSIONS:
        return _extract_ipynb_text(stream)
    if suffix in TEXT_EXTENSIONS or suffix in CODE_EXTENSIONS or suffix in DATA_EXTENSIONS:
        if suffix == ".json":
            try:
                parsed = json.loads(data.decode("utf-8", errors="ignore"))
                return json.dumps(parsed, indent=2, ensure_ascii=False), None
            except (json.JSONDecodeError, UnicodeDecodeError):
                pass
        return data.decode("utf-8", errors="ignore").strip(), None
    if suffix in MEDIA_EXTENSIONS:
        return "", "binary media file — no text extracted"
    try:
        text = data.decode("utf-8").strip()
        if text and sum(c.isprintable() or c in "\n\r\t" for c in text) / max(len(text), 1) > 0.85:
            return text, None
    except UnicodeDecodeError:
        pass
    return "", "binary file — no text extracted"


def _is_text_like(filename):
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_EXTENSIONS or suffix in CODE_EXTENSIONS or suffix in DATA_EXTENSIONS:
        return True
    if suffix in NOTEBOOK_EXTENSIONS:
        return True
    return False


def _file_kind(filename):
    """Return the technical kind used to choose a content viewer/editor."""
    suffix = Path(filename).suffix.lower()
    if suffix in (".md", ".markdown"):
        return "markdown"
    if suffix in (".csv",):
        return "csv"
    if suffix in (".tsv",):
        return "tsv"
    if suffix == ".json":
        return "json"
    if suffix == ".tex":
        return "tex"
    if suffix in NOTEBOOK_EXTENSIONS:
        return "ipynb"
    if suffix in SHEET_EXTENSIONS:
        return "spreadsheet"
    if suffix in SLIDE_EXTENSIONS:
        return "presentation"
    if suffix in CODE_EXTENSIONS:
        return "code"
    if suffix in TEXT_EXTENSIONS:
        return "text"
    if suffix in DOC_EXTENSIONS:
        return "document"
    if suffix in MEDIA_EXTENSIONS:
        return "media"
    return "binary"


def file_workspace_kind(filename):
    """Classify a file into the four repository landing-page tabs."""
    lower = Path(filename or "").name.lower()
    suffix = Path(lower).suffix
    # Waypoint-created slide/other files remain text-editable markdown while retaining
    # an explicit landing-page kind.
    if lower.endswith(".slides.md"):
        return "slides"
    if lower.endswith(".other.md"):
        return "other"
    if suffix in SHEET_EXTENSIONS:
        return "sheets"
    if suffix in SLIDE_EXTENSIONS:
        return "slides"
    if suffix in DOC_EXTENSIONS or suffix in TEXT_EXTENSIONS:
        return "docs"
    return "other"


def _guess_mime(filename):
    suffix = Path(filename).suffix.lower()
    if suffix in CODE_EXTENSIONS:
        return "text/plain"
    if suffix in TEXT_EXTENSIONS:
        return "text/plain"
    if suffix == ".json":
        return "application/json"
    if suffix == ".tex":
        return "text/x-tex"
    if suffix == ".ipynb":
        return "application/x-ipynb+json"
    if suffix == ".csv":
        return "text/csv"
    if suffix == ".tsv":
        return "text/tab-separated-values"
    if suffix == ".pdf":
        return "application/pdf"
    if suffix in (".docx", ".doc"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return "application/octet-stream"


def _extract_text(file_storage):
    """Return (text, note) where note is set for unreadable binaries."""
    filename = file_storage.filename or "upload"
    suffix = Path(filename).suffix.lower()
    stream = file_storage.stream
    stream.seek(0)
    try:
        if suffix == ".pdf":
            return _extract_pdf_text(stream).strip(), None
        if suffix in (".docx", ".doc"):
            return _extract_docx_text(stream).strip(), None
        if suffix in NOTEBOOK_EXTENSIONS:
            return _extract_ipynb_text(stream)
        if suffix in TEXT_EXTENSIONS or suffix in CODE_EXTENSIONS or suffix in DATA_EXTENSIONS:
            raw = stream.read()
            if suffix == ".json":
                try:
                    parsed = json.loads(raw.decode("utf-8", errors="ignore"))
                    return json.dumps(parsed, indent=2, ensure_ascii=False), None
                except (json.JSONDecodeError, UnicodeDecodeError):
                    pass
            return raw.decode("utf-8", errors="ignore").strip(), None
        if suffix in MEDIA_EXTENSIONS:
            return "", "binary media file — no text extracted"
        # Best-effort decode for unknown types
        raw = stream.read()
        try:
            text = raw.decode("utf-8").strip()
            if text and sum(c.isprintable() or c in "\n\r\t" for c in text) / max(len(text), 1) > 0.85:
                return text, None
        except UnicodeDecodeError:
            pass
        return "", "binary file — no text extracted"
    except Exception as exc:
        return "", f"extraction failed: {exc}"
    finally:
        stream.seek(0)


def _heuristic_category(filename, text):
    suffix = Path(filename).suffix.lower()
    lower_name = filename.lower()
    lower_text = (text or "").lower()

    if suffix in CODE_EXTENSIONS or any(k in lower_name for k in ("def ", "function", "import ", "class ")):
        if suffix in CODE_EXTENSIONS or re.search(r"\b(def |function |import |class )\b", lower_text):
            return "Code"
    if suffix == ".tex" or "overleaf" in lower_name or "\\begin{document}" in lower_text:
        return "Research"
    if suffix in DATA_EXTENSIONS or suffix in (".csv", ".tsv"):
        return "Data"
    if suffix in MEDIA_EXTENSIONS:
        return "Media"
    if suffix in DOC_EXTENSIONS:
        return "Docs"
    if suffix in TEXT_EXTENSIONS:
        if any(k in lower_text for k in ("todo", "journal", "note", "remember", "idea")):
            return "Notes"
        return "Notes" if suffix == ".md" else "Docs"

    keyword_map = {
        "Research": ("abstract", "hypothesis", "methodology", "bibliography", "citation", "experiment"),
        "Notes": ("todo", "journal", "reminder", "scratch", "memo"),
        "Data": ("dataset", "schema", "csv", "json", "table", "records"),
        "Code": ("def ", "function", "import ", "class ", "const ", "return ", "api"),
        "Docs": ("chapter", "section", "introduction", "overview", "documentation"),
    }
    scores = {cat: 0 for cat in CATEGORIES}
    for cat, keys in keyword_map.items():
        for key in keys:
            if key in lower_text or key in lower_name:
                scores[cat] += 1
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "Other"


def _ai_allowed():
    return db.is_api_enabled() and budget.cap_pressure() < AI_PRESSURE_MAX


def _ai_refine_category(filename, text, heuristic):
    if not _ai_allowed() or _word_count(text) < AI_MIN_WORDS:
        return heuristic
    sample = text[:2500]
    prompt = (
        f"Classify this file into exactly ONE category from: {', '.join(CATEGORIES)}.\n"
        f"Filename: {filename}\n"
        f"Heuristic guess: {heuristic}\n"
        f"Content preview:\n{sample}\n\n"
        'Return JSON: {"category": "<one of the fixed categories>"}'
    )
    try:
        result = ai.route_json("tasks", prompt, max_tokens=80)
        cat = (result.get("category") or "").strip()
        if cat in CATEGORIES:
            return cat
    except Exception as exc:
        print(f"[repository] AI category refine failed: {exc}")
    return heuristic


def _ai_file_summary(filename, category, text):
    if not _ai_allowed() or _word_count(text) < 15:
        words = _word_count(text)
        if words:
            preview = " ".join((text or "").split()[:25])
            return f"{filename} ({category}, {words} words): {preview}…"
        return f"{filename} ({category}): stored with no extractable text."
    sample = text[:2000]
    prompt = (
        f"Write a 1-2 sentence summary of this file for a personal knowledge repository.\n"
        f"Filename: {filename}\nCategory: {category}\n\n{sample}\n\n"
        'Return JSON: {"summary": "..."}'
    )
    try:
        result = ai._cheap_json(prompt, max_tokens=120)
        summary = (result.get("summary") or "").strip()
        if summary:
            return summary
    except Exception as exc:
        print(f"[repository] AI file summary failed: {exc}")
    preview = " ".join(text.split()[:30])
    return f"{filename}: {preview}…"


def _summary_stamp():
    return f"{db.repo_file_count()}:{db.repo_total_words()}"


def _bump_summary_staleness():
    settings = db.get_settings()
    if settings.get("repo_master_summary") or settings.get("repo_summary_stamp"):
        db.save_settings({"repo_summary_stamp": ""})


def _update_repo_file_fields(file_id, **fields):
    if not fields:
        return
    cols = ", ".join(f"{k} = ?" for k in fields)
    values = list(fields.values()) + [file_id]
    conn = db.get_conn()
    conn.execute(f"UPDATE repo_files SET {cols} WHERE id = ?", values)
    conn.commit()
    conn.close()


def _backup_file_version(file_id, stored_path):
    src = REPO_STORE / stored_path
    if not src.is_file():
        return
    dest_dir = VERSIONS_DIR / str(file_id)
    dest_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    shutil.copy2(src, dest_dir / ts)
    versions = sorted(dest_dir.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in versions[MAX_VERSION_BACKUPS:]:
        old.unlink(missing_ok=True)


def _python_executable():
    override = (os.environ.get("WAYPOINT_PYTHON") or "").strip()
    if override and Path(override).is_file():
        return override
    if VENV_PYTHON.is_file():
        return str(VENV_PYTHON)
    return sys.executable


def _invalidate_rag_cache():
    _rag_cache["file_count"] = None
    _rag_cache["vectorizer"] = None
    _rag_cache["matrix"] = None
    _rag_cache["chunks"] = []


def _chunk_text(text, file_id, filename, category):
    if not text:
        return [{"file_id": file_id, "filename": filename, "category": category, "text": "", "offset": 0}]
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        chunks.append({
            "file_id": file_id,
            "filename": filename,
            "category": category,
            "text": text[start:end],
            "offset": start,
        })
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def _build_rag_index():
    file_count = db.repo_file_count()
    if _rag_cache["file_count"] == file_count and _rag_cache["vectorizer"] is not None:
        return
    rows = db.list_repo_files()
    chunks = []
    for row in rows:
        chunks.extend(_chunk_text(
            row.get("extracted_text") or "",
            row["id"],
            row["filename"],
            row.get("category") or "Other",
        ))
    # Filter empty chunks for vectorization but keep metadata
    indexed = [c for c in chunks if (c["text"] or "").strip()]
    if not indexed:
        _rag_cache.update(file_count=file_count, vectorizer=None, matrix=None, chunks=[])
        return
    texts = [c["text"] for c in indexed]
    vectorizer = TfidfVectorizer(stop_words="english", max_features=4000)
    matrix = vectorizer.fit_transform(texts)
    _rag_cache.update(
        file_count=file_count,
        vectorizer=vectorizer,
        matrix=matrix,
        chunks=indexed,
    )


def _snippet_around(text, query_terms):
    if not text:
        return ""
    lower = text.lower()
    best_pos = 0
    for term in query_terms:
        pos = lower.find(term.lower())
        if pos >= 0:
            best_pos = pos
            break
    start = max(0, best_pos - SNIPPET_RADIUS)
    end = min(len(text), best_pos + SNIPPET_RADIUS)
    snippet = text[start:end].strip()
    if start > 0:
        snippet = "…" + snippet
    if end < len(text):
        snippet = snippet + "…"
    return snippet


def ingest_bytes(filename, data, *, category=None, mime=None):
    """Ingest raw bytes (Google import, zip extraction). Returns repo_files row."""
    suffix = Path(filename).suffix.lower()
    stored_name = f"{uuid.uuid4().hex}{suffix}"
    dest = REPO_STORE / stored_name
    dest.write_bytes(data)

    text, note = _extract_from_bytes(filename, data)
    if note and not text:
        text = f"[{note}]"

    mime = mime or _guess_mime(filename)
    size = len(data)
    words = _word_count(text)

    if category and category in CATEGORIES:
        chosen = category
    else:
        chosen = _heuristic_category(filename, text)
        chosen = _ai_refine_category(filename, text, chosen)
    summary = _ai_file_summary(filename, chosen, text)

    row = db.insert_repo_file(
        filename=filename,
        stored_path=stored_name,
        mime=mime,
        size=size,
        category=chosen,
        summary=summary,
        extracted_text=text,
        word_count=words,
    )
    _invalidate_rag_cache()
    _bump_summary_staleness()
    return row


def _ingest_zip_archive(file_storage):
    """Extract text-like files from a zip (Overleaf/LaTeX projects) into repo_files rows."""
    filename = file_storage.filename or "archive.zip"
    file_storage.stream.seek(0)
    data = file_storage.stream.read()
    rows = []
    skipped = 0
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                inner_name = Path(info.filename).name
                if not inner_name or inner_name.startswith("."):
                    skipped += 1
                    continue
                suffix = Path(inner_name).suffix.lower()
                if suffix not in ZIP_TEXT_EXTENSIONS:
                    skipped += 1
                    continue
                try:
                    inner_data = zf.read(info)
                except Exception:
                    skipped += 1
                    continue
                if not inner_data:
                    skipped += 1
                    continue
                category = "Research" if suffix == ".tex" else None
                rows.append(ingest_bytes(inner_name, inner_data, category=category))
    except zipfile.BadZipFile as exc:
        raise ValueError(f"invalid zip archive: {exc}") from exc
    if not rows:
        raise ValueError("zip contained no extractable text files")
    return {"rows": rows, "extracted": len(rows), "skipped": skipped}


def ingest_file(file_storage):
    """Save file bytes, extract text, categorize, summarize; return repo_files row."""
    filename = file_storage.filename or "upload"
    suffix = Path(filename).suffix.lower()
    if suffix == ".zip":
        return _ingest_zip_archive(file_storage)

    file_storage.stream.seek(0)
    data = file_storage.stream.read()
    return ingest_bytes(filename, data, mime=file_storage.mimetype or _guess_mime(filename))


def _safe_doc_filename(name, *, extracted=False):
    """Filesystem-safe markdown filename from a title or stem."""
    raw = (name or "untitled").strip()
    safe = re.sub(r"[^\w\s\-_.]", "", raw)
    safe = re.sub(r"\s+", "-", safe).strip("-_.") or "untitled"
    if extracted and not safe.lower().endswith(".extracted.md"):
        safe = f"{safe}.extracted" if safe.lower().endswith(".md") else f"{safe}.extracted.md"
    elif not safe.lower().endswith(".md"):
        safe = f"{safe}.md"
    return safe[:120]


def _managed_filename_suffix(filename):
    lower = (filename or "").lower()
    for suffix in (".slides.md", ".other.md"):
        if lower.endswith(suffix):
            return suffix
    return Path(filename or "").suffix


def rename_file(file_id, title):
    """Rename a repository item while preserving its format/kind suffix."""
    row = db.get_repo_file(file_id)
    if not row:
        return None, "not found"
    title = str(title or "").strip()
    if not title:
        return None, "title is required"
    if any(char in title for char in ("/", "\\", "\x00")):
        return None, "title cannot contain path separators"

    suffix = _managed_filename_suffix(row["filename"])
    if suffix and title.lower().endswith(suffix.lower()):
        title = title[:-len(suffix)].rstrip()
    safe = re.sub(r"[^\w\s\-().,'&]", "", title)
    safe = re.sub(r"\s+", " ", safe).strip(" .")
    if not safe:
        return None, "title must contain letters or numbers"
    filename = f"{safe[:120]}{suffix}"
    db.rename_repo_file(file_id, filename)
    _bump_summary_staleness()
    return db.get_repo_file(file_id), None


def get_file_content(file_id):
    """Return editable text content metadata for a stored file."""
    row = db.get_repo_file(file_id)
    if not row:
        return None
    kind = _file_kind(row["filename"])
    extracted = row.get("extracted_text") or ""
    extracted_truncated = len(extracted) > EXTRACTED_TEXT_RETURN_MAX
    if extracted_truncated:
        extracted = extracted[:EXTRACTED_TEXT_RETURN_MAX]
    payload = {
        "editable": False,
        "kind": kind,
        "extracted_text": extracted,
        "extracted_text_truncated": extracted_truncated,
        "extracted_text_total_chars": len(row.get("extracted_text") or ""),
        "word_count": row.get("word_count"),
        "filename": row["filename"],
        "category": row.get("category") or "Other",
    }
    if not _is_text_like(row["filename"]):
        return payload
    path = REPO_STORE / row["stored_path"]
    if not path.is_file():
        payload["error"] = "file missing on disk"
        return payload
    text = path.read_text(encoding="utf-8", errors="ignore")
    payload["text"] = text
    payload["editable"] = True
    return payload


def update_file_content(file_id, new_text):
    """Rewrite stored bytes, re-extract text, refresh summary, version backup."""
    row = db.get_repo_file(file_id)
    if not row:
        return None, "not found"
    if not _is_text_like(row["filename"]):
        return None, "file is not editable"
    path = REPO_STORE / row["stored_path"]
    if not path.is_file():
        return None, "file missing on disk"

    _backup_file_version(file_id, row["stored_path"])
    data = (new_text or "").encode("utf-8")
    path.write_bytes(data)

    text, note = _extract_from_bytes(row["filename"], data)
    if note and not text:
        text = f"[{note}]"
    words = _word_count(text)
    category = row.get("category") or "Other"
    summary = _ai_file_summary(row["filename"], category, text)

    _update_repo_file_fields(
        file_id,
        extracted_text=text,
        word_count=words,
        summary=summary,
        size=len(data),
    )
    _invalidate_rag_cache()
    _bump_summary_staleness()
    return db.get_repo_file(file_id), None


def run_python_file(file_id):
    """Execute a stored .py file locally. SECURITY: explicit user-triggered only."""
    row = db.get_repo_file(file_id)
    if not row:
        return None, "not found"
    suffix = Path(row["filename"]).suffix.lower()
    if suffix != ".py":
        return None, "only .py files can be run"
    path = REPO_STORE / row["stored_path"]
    if not path.is_file():
        return None, "file missing on disk"
    try:
        proc = proc_util.run(
            [_python_executable(), str(path)],
            cwd=str(REPO_STORE),
            capture_output=True,
            text=True,
            timeout=RUN_TIMEOUT_SECONDS,
        )
        return {
            "stdout": proc.stdout or "",
            "stderr": proc.stderr or "",
            "exit_code": proc.returncode,
        }, None
    except subprocess.TimeoutExpired:
        return {
            "stdout": "",
            "stderr": f"Timed out after {RUN_TIMEOUT_SECONDS}s",
            "exit_code": -1,
        }, None
    except Exception as exc:
        return None, str(exc)


def reextract_file(file_id):
    """Re-run text extraction from stored bytes and refresh summary."""
    row = db.get_repo_file(file_id)
    if not row:
        return None, "not found"
    path = REPO_STORE / row["stored_path"]
    if not path.is_file():
        return None, "file missing on disk"

    data = path.read_bytes()
    text, note = _extract_from_bytes(row["filename"], data)
    if note and not text:
        text = f"[{note}]"
    words = _word_count(text)
    category = row.get("category") or "Other"
    summary = _ai_file_summary(row["filename"], category, text)

    _update_repo_file_fields(
        file_id,
        extracted_text=text,
        word_count=words,
        summary=summary,
    )
    _invalidate_rag_cache()
    _bump_summary_staleness()
    return db.get_repo_file(file_id), None


def promote_extracted_to_docs(file_id, title=None):
    """Create a new Docs markdown file from a file's extracted_text."""
    row = db.get_repo_file(file_id)
    if not row:
        return None, "not found"
    text = (row.get("extracted_text") or "").strip()
    if not text or text.startswith("["):
        return None, "no extractable text to promote"

    if title and str(title).strip():
        filename = _safe_doc_filename(str(title).strip())
    else:
        stem = Path(row["filename"]).stem
        filename = _safe_doc_filename(stem, extracted=True)

    new_row = ingest_bytes(filename, text.encode("utf-8"), category="Docs")
    return new_row, None


def create_blank_doc(title="Untitled"):
    """Create an empty/minimal markdown Docs file."""
    return create_blank_file(title=title, kind="docs")


def create_blank_file(title="Untitled", kind="docs"):
    """Create a locally editable file for one repository landing-page kind."""
    kind = str(kind or "docs").strip().lower()
    if kind not in WORKSPACE_KINDS:
        return None, f"kind must be one of: {', '.join(WORKSPACE_KINDS)}"
    title = str(title or "Untitled").strip() or "Untitled"
    safe_title = re.sub(r"[^\w\s\-_.]", "", title)
    safe_title = re.sub(r"\s+", "-", safe_title).strip("-_.") or "untitled"
    safe_title = safe_title[:100]
    specs = {
        "docs": (".md", "Docs", f"# {title}\n\n"),
        "sheets": (".csv", "Data", "Column A,Column B,Column C\n,,\n"),
        "slides": (".slides.md", "Docs", f"# {title}\n\n## Slide 1\n\n"),
        "other": (".other.md", "Other", f"# {title}\n\n"),
    }
    suffix, category, content = specs[kind]
    filename = safe_title if safe_title.lower().endswith(suffix) else f"{safe_title}{suffix}"
    return ingest_bytes(filename, content.encode("utf-8"), category=category), None


def ask_docs(query, k=6, file_id=None):
    """RAG search + optional AI answer grounded in repository excerpts."""
    query = (query or "").strip()
    if not query:
        return {"answer": "Please provide a question.", "results": [], "grounded": False}

    try:
        k = int(k)
    except (TypeError, ValueError):
        k = 6
    k = max(1, min(k, 30))

    results = rag_search(query, k=k)
    scoped_file = None
    if file_id is not None:
        try:
            fid = int(file_id)
            results = [r for r in results if r["file_id"] == fid]
            scoped_file = db.get_repo_file(fid)
        except (TypeError, ValueError):
            pass

    if not results and not (scoped_file and (scoped_file.get("extracted_text") or "").strip()):
        return {
            "answer": "No matching documents found in the repository.",
            "results": [],
            "grounded": False,
        }

    context_parts = []
    total = 0
    for r in results:
        # Prefer full RAG chunk text — short display snippets were only the start of each hit.
        body = (r.get("chunk_text") or r.get("snippet") or "").strip()
        if not body:
            continue
        block = f"--- {r['filename']} (category: {r['category']})\n{body}\n"
        if total + len(block) > ASK_DOCS_CONTEXT_CHARS:
            remaining = ASK_DOCS_CONTEXT_CHARS - total
            if remaining > 200:
                context_parts.append(block[:remaining] + "…")
            break
        context_parts.append(block)
        total += len(block)

    # Single-file ask: if RAG found little, feed more of the document (not just the opening).
    if scoped_file and total < min(8000, ASK_DOCS_CONTEXT_CHARS):
        full = (scoped_file.get("extracted_text") or "").strip()
        if full:
            remaining = ASK_DOCS_CONTEXT_CHARS - total
            if remaining > 400:
                name = scoped_file.get("filename") or "document"
                cat = scoped_file.get("category") or "Docs"
                # Prefer a mid-document window when the query matches later terms.
                terms = [t.lower() for t in re.findall(r"\w+", query) if len(t) > 2]
                start = 0
                if terms:
                    lower = full.lower()
                    hits = [lower.find(t) for t in terms if lower.find(t) >= 0]
                    if hits:
                        start = max(0, min(hits) - 800)
                window = full[start : start + remaining]
                prefix = "…" if start > 0 else ""
                suffix = "…" if start + len(window) < len(full) else ""
                context_parts.append(
                    f"--- {name} (category: {cat}) [document window]\n{prefix}{window}{suffix}\n"
                )
                total += len(window)

    context = "\n".join(context_parts)

    if _ai_allowed():
        prompt = (
            "Answer the user's question using ONLY the repository excerpts below. "
            "If the excerpts do not contain enough information, say so clearly. "
            "Cite source filenames in parentheses when referencing specific content.\n\n"
            f"Question: {query}\n\n"
            f"Excerpts:\n{context}"
        )
        try:
            answer = ai.route_text("tasks", prompt, max_tokens=ASK_DOCS_MAX_TOKENS)
            return {"answer": answer, "results": results, "grounded": True}
        except Exception as exc:
            print(f"[repository] ask_docs AI failed: {exc}")

    snippets = [r["snippet"] for r in results[:3] if r.get("snippet")]
    answer = " ".join(snippets) if snippets else "Matching excerpts found but AI is unavailable."
    if len(results) > 3:
        answer += f" … (+{len(results) - 3} more matches)"
    return {"answer": answer, "results": results, "grounded": False}


def propose_planning_from_file(file_id, intent="tasks"):
    """Extract file text and return planning proposals only (no DB writes)."""
    row = db.get_repo_file(file_id)
    if not row:
        return None, "not found"
    text = (row.get("extracted_text") or "").strip()
    if not text or text.startswith("["):
        return None, "no extractable text in this file"
    intent = (intent or "tasks").strip().lower()
    proposals = {"tasks": [], "schedule": [], "proposal_only": True, "intent": intent}

    if intent in ("tasks", "auto", "all"):
        try:
            proposals["tasks"] = orchestrator.run("freeform_task_parser", text=text) or []
        except Exception as exc:
            proposals["tasks_error"] = str(exc)

    if intent in ("schedule", "days", "auto", "all"):
        try:
            proposals["schedule"] = orchestrator.run(
                "class_schedule_parser", text=text, today=date_cls.today().isoformat(),
            ) or []
        except Exception as exc:
            proposals["schedule_error"] = str(exc)

    return proposals, None


def import_google_drive_files(max_files=50):
    """Copy Google Drive files locally (never mutates originals)."""
    import google_integration

    if not google_integration.is_connected():
        return {
            "ok": False,
            "needs_auth": True,
            "how": "Connect Google in Settings → Integrations",
        }

    cap = max(1, min(int(max_files or 50), 100))
    imported = []
    skipped = 0
    errors = []
    scope_note = google_integration.drive_scope_status()

    for item in google_integration.iter_drive_files_for_import(max_results=cap):
        name = item.get("name") or "untitled"
        try:
            filename, data, mime = google_integration.export_drive_file_copy(item)
            if not data:
                skipped += 1
                continue
            row = ingest_bytes(filename, data, category="Google", mime=mime)
            imported.append({"file_id": row["id"], "filename": row["filename"], "source": name})
        except Exception as exc:
            errors.append({"name": name, "error": str(exc)})
            skipped += 1

    return {
        "ok": True,
        "imported": len(imported),
        "skipped": skipped,
        "files": imported,
        "errors": errors[:20],
        "scope_note": scope_note,
    }


def delete_file(file_id):
    row = db.get_repo_file(file_id)
    if not row:
        return False
    stored = REPO_STORE / row["stored_path"]
    if stored.is_file():
        stored.unlink()
    db.delete_repo_file(file_id)
    _invalidate_rag_cache()
    return True


def rag_search(query, k=8):
    query = (query or "").strip()
    if not query:
        return []
    try:
        k = int(k)
    except (TypeError, ValueError):
        k = 8
    k = max(1, min(k, 30))

    _build_rag_index()
    if not _rag_cache["chunks"] or _rag_cache["vectorizer"] is None:
        return []

    vectorizer = _rag_cache["vectorizer"]
    matrix = _rag_cache["matrix"]
    chunks = _rag_cache["chunks"]

    q_vec = vectorizer.transform([query])
    scores = cosine_similarity(q_vec, matrix)[0]
    top_idx = np.argsort(scores)[::-1][:k]

    query_terms = [t for t in re.findall(r"\w+", query) if len(t) > 2]
    results = []
    seen_files = set()
    for idx in top_idx:
        score = float(scores[idx])
        if score <= 0:
            continue
        chunk = chunks[idx]
        if chunk["file_id"] in seen_files and len(results) >= k:
            continue
        seen_files.add(chunk["file_id"])
        results.append({
            "file_id": chunk["file_id"],
            "filename": chunk["filename"],
            "category": chunk["category"],
            "score": round(score, 4),
            "snippet": _snippet_around(chunk["text"], query_terms),
            "chunk_text": chunk["text"],
            "offset": chunk["offset"],
        })
    return results


def rebuild_summary(force=False):
    stamp = _summary_stamp()
    settings = db.get_settings()
    if not force and settings.get("repo_summary_stamp") == stamp and settings.get("repo_master_summary"):
        return {
            "summary": settings["repo_master_summary"],
            "categories": list_repo_category_counts(),
            "total_files": db.repo_file_count(),
            "total_words": db.repo_total_words(),
            "generated_at": settings.get("repo_summary_generated_at"),
            "cached": True,
        }

    files = db.list_repo_files()
    categories = list_repo_category_counts()
    total_files = len(files)
    total_words = sum(f.get("word_count") or 0 for f in files)

    if not files:
        summary = "Repository is empty. Upload files to build your local knowledge dump."
    elif _ai_allowed():
        lines = []
        for f in files[:80]:
            lines.append(
                f"- {f['filename']} [{f.get('category', 'Other')}]: "
                f"{(f.get('summary') or '')[:200]}"
            )
        if total_files > 80:
            lines.append(f"... and {total_files - 80} more files")
        cat_line = ", ".join(f"{c['name']} ({c['count']})" for c in categories)
        prompt = (
            "You are summarizing a personal file repository. Produce a structured rolling overview "
            "(3-6 short paragraphs or bullet sections) covering themes, key topics, and how content "
            "is distributed. Be concise.\n\n"
            f"Total files: {total_files}, total words: {total_words}\n"
            f"Categories: {cat_line}\n\n"
            "Per-file summaries:\n" + "\n".join(lines) + "\n\n"
            'Return JSON: {"summary": "..."}'
        )
        try:
            result = ai._cheap_json(prompt, max_tokens=600)
            summary = (result.get("summary") or "").strip()
        except Exception as exc:
            print(f"[repository] master summary AI failed: {exc}")
            summary = _fallback_master_summary(files, categories, total_files, total_words)
    else:
        summary = _fallback_master_summary(files, categories, total_files, total_words)

    generated_at = db.now_iso()
    db.save_settings({
        "repo_master_summary": summary,
        "repo_summary_stamp": stamp,
        "repo_summary_generated_at": generated_at,
    })
    return {
        "summary": summary,
        "categories": categories,
        "total_files": total_files,
        "total_words": total_words,
        "generated_at": generated_at,
        "cached": False,
    }


def _fallback_master_summary(files, categories, total_files, total_words):
    cat_line = ", ".join(f"{c['name']}: {c['count']}" for c in categories) or "none"
    names = ", ".join(f["filename"] for f in files[:12])
    if total_files > 12:
        names += f", … (+{total_files - 12} more)"
    return (
        f"Repository holds {total_files} files ({total_words:,} words). "
        f"Categories: {cat_line}. Files include: {names}."
    )


def list_repo_category_counts():
    return db.list_repo_categories()


def stored_file_path(file_id):
    row = db.get_repo_file(file_id)
    if not row:
        return None
    return REPO_STORE / row["stored_path"]
